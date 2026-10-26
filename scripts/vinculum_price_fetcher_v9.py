#!/usr/bin/env python3
"""
Vinculum Protocol — Universal Price Fetcher v9
===============================================
Reference price lookup for vault mint ratio calculations.
Run twice per day (VF-ORC-001). Target: within a few percent. Not for trading.
First valid cascade source wins per asset; later sources are not consulted (VF-ORC-002).
No two-source/spread/median/TWAP/10-min freshness extras (VF-ORC-003).
No hardcoded substitute prices (VF-ORC-004). Miss => asset unavailable until a later run (VF-ORC-005).

Built against: approved_assets_final.json (1001 assets, clean)

Cascade per asset (Revision 18 Section 7, D-64, VF-ORC-006):
  Tier 1: CoinGecko  (batch 200 IDs per call)
  Tier 2: DexScreener (the registry's own contract address on its own chain)

No symbol search (removed by D-64) and no community-token override
(VF-ORC-006: no override applies, and a displayed registry contract is never
replaced with a different lookup contract). The former v9 overrides for
TigerOG, LionOG, FrogOG (BSC legacy contracts) and WKC are removed under
t427u. TigerOG is on the Base registry by founder override (t430u, t431u).
The override never supplies a price. Under operator decision t436u its Tier 2
lookup takes a DexScreener pool for its exact registry contract on any chain
(TigerOG is an Axelar ITS token, so the same address exists on other chains),
the pool with the most liquidity, and no market-activity, trading-recency or
pool-age test applies to it. Every other row uses its own chain's pools only.
No hardcoded prices anywhere.

Usage:
  python vinculum_price_fetcher_v9.py

Both files must be in the same folder:
  approved_assets_final.json
  vinculum_price_fetcher_v9.py

Outputs:
  vinculum_prices.json
  vinculum_prices.xlsx
  price_fetch_log.txt
"""

import json, os, sys, time
from datetime import datetime, timezone

try:
    import requests
except ImportError:
    sys.exit("Run: pip install requests openpyxl")

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

# ── CONFIG ───────────────────────────────────────────────────────────────────
INPUT_FILE  = "approved_assets_final.json"
OUTPUT_JSON = "vinculum_prices.json"
OUTPUT_XLSX = "vinculum_prices.xlsx"
LOG_FILE    = "price_fetch_log.txt"

TIMEOUT_S   = 14
CG_BATCH    = 200
CG_PAUSE    = 4.0   # seconds between CoinGecko batch calls
DS_PAUSE    = 0.6   # DexScreener
MAX_RETRY   = 3

NATIVE_CHAINS = {
    "Bitcoin","Dogecoin","Litecoin","DigiByte","Zcash","XRPL","Stellar","Cosmos"
}

DS_CHAIN = {
    "Ethereum":"ethereum","Base":"base","Arbitrum":"arbitrum",
    "Optimism":"optimism","Polygon":"polygon","BSC":"bsc",
    "Avalanche":"avalanche","Solana":"solana",
}

fh = None

def log(msg):
    ts = datetime.now().strftime("%H:%M:%S")
    line = f"[{ts}] {msg}"
    print(line)
    if fh and not fh.closed:
        fh.write(line + "\n"); fh.flush()

def now_ts():
    return datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")

def fetch(url, params=None):
    for attempt in range(1, MAX_RETRY + 1):
        try:
            r = requests.get(url, params=params, timeout=TIMEOUT_S,
                             headers={"Accept":"application/json",
                                      "User-Agent":"vinculum-price-fetcher/9"})
            if r.status_code == 200:
                return r.json()
            if r.status_code == 429:
                wait = 12 * attempt
                log(f"  429 rate-limit -- waiting {wait}s")
                time.sleep(wait); continue
            if r.status_code == 404:
                return None
        except Exception as e:
            log(f"  fetch error (attempt {attempt}): {e}")
        time.sleep(3 * attempt)
    return None

# ── TIER 1: COINGECKO BATCH ──────────────────────────────────────────────────
def cg_batch(id_list):
    """Fetch prices for a list of CoinGecko IDs. Returns {cg_id: usd}."""
    result = {}
    for i in range(0, len(id_list), CG_BATCH):
        batch = id_list[i:i+CG_BATCH]
        data = fetch("https://api.coingecko.com/api/v3/simple/price",
                     params={"ids": ",".join(batch), "vs_currencies": "usd"})
        if data:
            for cg_id, vals in data.items():
                if vals.get("usd"):
                    result[cg_id] = float(vals["usd"])
        if i + CG_BATCH < len(id_list):
            time.sleep(CG_PAUSE)
    return result

# ── TIER 2: DEXSCREENER BY CONTRACT ──────────────────────────────────────────
# The price is always the price of the token being priced. A DexScreener pair
# reports priceUsd for its base (first) token only. When the token is the base
# token, priceUsd is its price. When it is the quote (second) token, its price
# is priceUsd / priceNative (priceNative is the base token's price in quote
# tokens). A pair that does not hold the token on either side, or lacks the
# fields needed, is skipped. Any pool may give the reference price; the pair
# with the most liquidity among the usable ones is taken. (Operator t431u: the
# pool choice does not matter; the other token's price is never assigned.)
def _f(x):
    try:
        return float(x or 0)
    except (TypeError, ValueError):
        return 0.0

def _same(a, b):
    return isinstance(a, str) and isinstance(b, str) and a.lower() == b.lower()

def token_usd_from_pair(pair, contract):
    """USD price of `contract` from one DexScreener pair, or None."""
    if not pair or not contract:
        return None
    price_usd = _f(pair.get("priceUsd"))
    if price_usd <= 0:
        return None
    if _same((pair.get("baseToken") or {}).get("address"), contract):
        return price_usd
    if _same((pair.get("quoteToken") or {}).get("address"), contract):
        price_native = _f(pair.get("priceNative"))
        if price_native <= 0:
            return None
        usd = price_usd / price_native
        return usd if usd > 0 and usd != float("inf") else None
    return None

# Founder override rows (operator t430u item 5, narrowed to TigerOG by t431u).
# Exact identity: registry environment and exact contract, never a ticker.
FOUNDER_OVERRIDE = (("Base", "0xCF7Fc0De71238c9EC45EC2Fd24FDc8521345dbB5"),)

def is_founder_override(chain, contract):
    return any(chain == c and _same(contract, a) for c, a in FOUNDER_OVERRIDE)

def best_ds_price(pairs, ds_chain, contract, any_chain=False):
    """Price from the usable pair with the most liquidity. Only pairs on
    ds_chain count, unless any_chain is set (founder override rows only).
    No volume, trade-count, recency or pool-age test is applied."""
    best = None
    for p in pairs or []:
        if not p or (not any_chain and p.get("chainId") != ds_chain):
            continue
        usd = token_usd_from_pair(p, contract)
        if usd is None:
            continue
        liq = _f((p.get("liquidity") or {}).get("usd"))
        if best is None or liq > best[0]:
            best = (liq, usd)
    return None if best is None else best[1]

def ds_by_contract(chain, contract):
    ds_chain = DS_CHAIN.get(chain)
    if not ds_chain or not contract:
        return None
    # Skip placeholder values
    if any(contract.upper().startswith(x) for x in
           ("MISSING","PENDING","REVIEW","NATIVE","DEADBEEF")):
        return None
    # Chain-scoped endpoint first: it lists the token's pools on its own chain.
    # The older all-chain endpoint returns a capped list that can leave out the
    # token's own-chain pools, so it is only a fallback.
    scoped = fetch(f"https://api.dexscreener.com/token-pairs/v1/{ds_chain}/{contract}")
    scoped = scoped if isinstance(scoped, list) else []
    if is_founder_override(chain, contract):
        # Operator t436u: any pool for this exact contract on any chain, the
        # most liquid one, with no market-activity requirement.
        data = fetch(f"https://api.dexscreener.com/latest/dex/tokens/{contract}")
        every = scoped + ((data or {}).get("pairs") or [])
        return best_ds_price(every, ds_chain, contract, any_chain=True)
    usd = best_ds_price(scoped, ds_chain, contract)
    if usd is not None:
        return usd
    data = fetch(f"https://api.dexscreener.com/latest/dex/tokens/{contract}")
    if not data or not data.get("pairs"):
        return None
    return best_ds_price(data["pairs"], ds_chain, contract)

# ── FORMAT ───────────────────────────────────────────────────────────────────
def fmt(usd):
    if usd is None: return "--"
    if usd >= 1000:       return f"${usd:,.2f}"
    if usd >= 1:          return f"${usd:.4f}"
    if usd >= 0.0001:     return f"${usd:.6f}"
    if usd >= 0.000000001: return f"${usd:.12f}"
    return f"${usd:.2e}"

# ── MAIN ─────────────────────────────────────────────────────────────────────
def main():
    global fh

    if not os.path.exists(INPUT_FILE):
        sys.exit(f"'{INPUT_FILE}' not found in current folder.")

    assets = json.load(open(INPUT_FILE, encoding="utf-8"))
    fh = open(LOG_FILE, "w", encoding="utf-8")
    ts  = now_ts()

    log("=== Vinculum Price Fetcher v9 ===")
    log(f"Assets: {len(assets)} | Purpose: vault reference rates (2x/day)")
    log(f"Input:  {INPUT_FILE}")
    log("")

    prices = {}

    # ── STEP 1: COINGECKO BATCH ──────────────────────────────────────────────
    log("Step 1: CoinGecko batch...")
    cg_map = {}
    for a in assets:
        cgid = a.get("coingecko_id","")
        if cgid:
            cg_map.setdefault(cgid, []).append((a["symbol"], a["chain"]))

    log(f"  {len(cg_map)} unique CoinGecko IDs across {len(assets)} assets")
    cg_prices = cg_batch(list(cg_map.keys()))
    applied = 0
    for cgid, asset_list in cg_map.items():
        usd = cg_prices.get(cgid)
        if usd:
            for sym, chain in asset_list:
                key = f"{sym}@{chain}"
                if key not in prices:
                    prices[key] = {"usd": usd, "source": "CoinGecko", "ts": ts}
                    applied += 1
    missed = len(cg_map) - len(cg_prices)
    log(f"  Got {len(cg_prices)} prices -> {applied} assets covered ({missed} IDs returned no price)")

    # ── STEP 2: DEXSCREENER BY CONTRACT (for batch misses) ───────────────────
    log("Step 2: DexScreener (contract address)...")
    ds_needed = [a for a in assets
                 if f"{a['symbol']}@{a['chain']}" not in prices
                 and a["chain"] in DS_CHAIN
                 and a.get("contract","")]
    log(f"  {len(ds_needed)} assets to try...")
    ds_applied = 0
    for a in ds_needed:
        key = f"{a['symbol']}@{a['chain']}"
        usd = ds_by_contract(a["chain"], a.get("contract",""))
        if usd:
            prices[key] = {"usd": usd, "source": "DexScreener", "ts": ts}
            ds_applied += 1
        time.sleep(DS_PAUSE)
    log(f"  DexScreener contract: {ds_applied} prices found")

    # ── STEP 3: MARK REMAINING ────────────────────────────────────────────────
    no_price = []
    for a in assets:
        key = f"{a['symbol']}@{a['chain']}"
        if key not in prices:
            prices[key] = {"usd": None, "source": "none", "ts": ts,
                           "error": "no price source resolved"}
            no_price.append(key)

    # ── SUMMARY ──────────────────────────────────────────────────────────────
    total  = len(assets)
    priced = sum(1 for a in assets
                 if prices.get(f"{a['symbol']}@{a['chain']}",{}).get("usd") is not None)
    from collections import Counter
    sources = Counter(prices[f"{a['symbol']}@{a['chain']}"]["source"]
                      for a in assets
                      if prices.get(f"{a['symbol']}@{a['chain']}",{}).get("usd"))

    log("")
    log("=== FINAL SUMMARY ===")
    log(f"Total assets:  {total}")
    log(f"Priced:        {priced} ({100*priced/total:.1f}%)")
    log(f"No price:      {total - priced}")
    log("")
    for src, cnt in sources.most_common():
        log(f"  {src:30} {cnt}")
    if no_price:
        log(f"\nNot priced ({len(no_price)}):")
        for k in no_price:
            log(f"  {k}")

    # ── SAVE JSON ─────────────────────────────────────────────────────────────
    with open(OUTPUT_JSON, "w", encoding="utf-8") as f:
        json.dump(prices, f, indent=1)
    log(f"\nJSON saved: {OUTPUT_JSON}")

    # ── SAVE XLSX ─────────────────────────────────────────────────────────────
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter

        wb = Workbook()
        ws = wb.active
        ws.title = "Prices"
        headers = ["Symbol","Chain","Price (USD)","Source","Timestamp",
                   "CoinGecko ID","Contract","Note"]
        ws.append(headers)
        hfill = PatternFill("solid", start_color="1A1A2E")
        hfont = Font(name="Arial", bold=True, color="FFFFFF", size=10)
        for c in ws[1]:
            c.fill = hfill; c.font = hfont
            c.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 22

        src_fills = {
            "CoinGecko":          PatternFill("solid", start_color="E3F2FD"),
            "DexScreener":        PatternFill("solid", start_color="EDE7F6"),
            "GeckoTerminal":      PatternFill("solid", start_color="FFF3E0"),
            "none":               PatternFill("solid", start_color="FFEBEE"),
        }

        for a in assets:
            key = f"{a['symbol']}@{a['chain']}"
            val = prices.get(key, {})
            usd = val.get("usd")
            ws.append([
                a["symbol"], a["chain"], fmt(usd),
                val.get("source",""), val.get("ts",""),
                a.get("coingecko_id",""), a.get("contract",""),
                val.get("error","")
            ])
            r = ws.max_row
            sfill = src_fills.get(val.get("source","none"),
                                   PatternFill("solid", start_color="FFFFFF"))
            for col in range(1, 9):
                ws.cell(row=r, column=col).font = Font(name="Arial", size=9)
            ws.cell(row=r, column=3).alignment = Alignment(horizontal="right")
            ws.cell(row=r, column=1).fill = sfill
            ws.cell(row=r, column=4).fill = sfill

        for i, w in enumerate([14, 11, 16, 20, 20, 28, 46, 36], 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"

        # Summary sheet
        ss = wb.create_sheet("Summary")
        ss.append(["Vinculum Price Fetcher v9 -- Run Summary"])
        ss.append(["Run time", ts])
        ss.append(["Input file", INPUT_FILE])
        ss.append(["Total assets", total])
        ss.append(["Priced", priced])
        ss.append(["Coverage", f"{100*priced/total:.1f}%"])
        ss.append([])
        ss.append(["Source", "Count"])
        for src, cnt in sources.most_common():
            ss.append([src, cnt])
        if no_price:
            ss.append([])
            ss.append(["Not priced", ""])
            for k in no_price:
                ss.append(["", k])
        ss["A1"].font = Font(name="Arial", bold=True, size=12)
        ss.column_dimensions["A"].width = 30
        ss.column_dimensions["B"].width = 50

        wb.save(OUTPUT_XLSX)
        log(f"Excel saved: {OUTPUT_XLSX}")

    except ImportError:
        log("openpyxl not installed -- xlsx skipped. JSON output is complete.")

    log(f"\nDone. {priced}/{total} assets priced ({100*priced/total:.1f}% coverage).")
    fh.close()


if __name__ == "__main__":
    main()